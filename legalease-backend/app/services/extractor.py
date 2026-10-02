from pathlib import Path
import io
import fitz
from docx import Document as DocxDocument
from PIL import Image
import pytesseract


def _extract_pdf_text(path: str) -> str:
    doc = fitz.open(path)
    try:
        pages = [page.get_text("text") for page in doc]
        text = "\n\n".join(pages).strip()
        if text:
            return text

        # Scanned/image PDF: OCR each page.
        ocr_pages = []
        for page in doc:
            pix = page.get_pixmap(matrix=fitz.Matrix(1.7, 1.7), alpha=False)
            image = Image.open(io.BytesIO(pix.tobytes("png")))
            ocr_pages.append(pytesseract.image_to_string(image))
        return "\n\n".join(ocr_pages).strip()
    finally:
        doc.close()


def _extract_docx_text(path: str) -> str:
    doc = DocxDocument(path)
    chunks = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                chunks.append(" | ".join(values))
    return "\n".join(chunks).strip()


def extract_text(path: str, filename: str) -> tuple[str, bool]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        text = _extract_pdf_text(path)
        return text, False if text else True
    if suffix == ".docx":
        return _extract_docx_text(path), False
    raise ValueError("Only PDF and DOCX files are supported")
